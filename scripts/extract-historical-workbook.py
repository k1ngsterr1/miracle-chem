#!/usr/bin/env python3
"""Read the two supplied 2026 workbooks without changing them; emit audited JSON.

Run with a Python environment containing openpyxl. The output retains every
non-empty cell, its cached value and its Excel formula when present. Parsed
records are read-only historical facts, never stock or accounting postings.
"""

import argparse
import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


MONTHS = {"июля": 7, "августа": 8, "сентября": 9, "октября": 10}


def value(cell):
    item = cell.value
    if isinstance(item, (datetime, date)):
        return item.date().isoformat() if isinstance(item, datetime) else item.isoformat()
    if isinstance(item, (int, float)) and not isinstance(item, bool):
        return str(Decimal(str(item)))
    return str(item).strip() if item is not None else None


def decimal(cell, *, allow_none=False):
    raw = value(cell)
    if raw is None and allow_none:
        return None
    if raw is None:
        raise ValueError(f"Missing number at {cell.coordinate}")
    try:
        return Decimal(raw.replace(",", "."))
    except Exception as exc:
        raise ValueError(f"Invalid number at {cell.coordinate}: {raw}") from exc


def number(cell, *, allow_none=False):
    item = decimal(cell, allow_none=allow_none)
    return format(item.normalize(), "f") if item is not None else None


def parse_date(raw):
    if raw is None:
        return None
    if isinstance(raw, (datetime, date)):
        return raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat()
    text = str(raw).strip().lower()
    match = re.fullmatch(r"(\d{1,2})\s+(июля|августа|сентября|октября)", text)
    if match:
        return date(2026, MONTHS[match.group(2)], int(match.group(1))).isoformat()
    if re.fullmatch(r"2026-\d\d-\d\d", text):
        return text
    return None


def source_rows(sheet, start, end):
    return [i for i in range(start, end + 1) if any(sheet.cell(i, c).value is not None for c in range(1, sheet.max_column + 1))]


def ingredient(sheet, row, name_col, rate_col, volume_col, used_col):
    name = value(sheet.cell(row, name_col))
    if not name:
        raise ValueError(f"Missing ingredient at {sheet.title}!{row}")
    rate = decimal(sheet.cell(row, rate_col))
    volume = decimal(sheet.cell(row, volume_col))
    consumed = decimal(sheet.cell(row, used_col))
    if abs(rate * volume - consumed) > Decimal("0.00001"):
        raise ValueError(f"Ingredient calculation mismatch at {sheet.title}!{row}")
    return {"name": name, "rate": format(rate.normalize(), "f"), "consumed": format(consumed.normalize(), "f"), "sourceRow": row}


def september_production(sheet):
    records = []
    for row in range(1, sheet.max_row + 1):
        day = parse_date(sheet.cell(row, 2).value)
        if day is None:
            continue
        start = row + 1
        end = start
        while end <= sheet.max_row and value(sheet.cell(end, 3)):
            end += 1
        end -= 1
        if end < start:
            raise ValueError(f"Production date without ingredients: {day}, row {row}")
        customer = value(sheet.cell(start, 2))
        if not customer or customer.isdigit():
            raise ValueError(f"Missing customer at production row {start}")
        volume = number(sheet.cell(start, 5))
        items = [ingredient(sheet, i, 3, 4, 5, 6) for i in range(start, end + 1)]
        if any(number(sheet.cell(i, 5)) != volume for i in range(start, end + 1)):
            raise ValueError(f"Inconsistent production volume on {day}")
        records.append({"sourceSheet": sheet.title, "sourceRows": source_rows(sheet, row, end),
                        "dateFrom": day, "dateTo": day, "customer": customer, "volumeM3": volume,
                        "ingredients": items})
    return records


def summer_production(sheet):
    groups = [
        (5, 9, "2026-07-28", "2026-07-28", "1", 3, 4, 5, 6),
        (11, 15, "2026-07-28", "2026-07-28", "2", 3, 4, 5, 6),
        (18, 22, "2026-07-29", "2026-07-29", None, 3, 4, 5, 6),
        (25, 29, "2026-07-31", "2026-07-31", None, 3, 4, 5, 6),
        (32, 36, "2026-08-01", "2026-08-01", None, 3, 4, 5, 6),
        (40, 44, "2026-08-02", "2026-08-02", None, 3, 4, 5, 6),
        (49, 53, "2026-08-04", "2026-08-14", None, 2, 6, 5, 7),
        (57, 61, "2026-08-17", "2026-08-17", None, 2, 6, 5, 7),
        (64, 68, "2026-08-21", "2026-08-21", None, 2, 6, 5, 7),
    ]
    records = []
    for start, end, from_date, to_date, part, name_col, rate_col, volume_col, used_col in groups:
        items = [ingredient(sheet, i, name_col, rate_col, volume_col, used_col) for i in range(start, end + 1)]
        volume = number(sheet.cell(start, volume_col))
        if any(number(sheet.cell(i, volume_col)) != volume for i in range(start, end + 1)):
            raise ValueError(f"Inconsistent production volume at row {start}")
        records.append({"sourceSheet": sheet.title, "sourceRows": source_rows(sheet, start, end),
                        "dateFrom": from_date, "dateTo": to_date, "part": part, "customer": None,
                        "volumeM3": volume, "ingredients": items})
    return records


def formulas(sheet):
    starts = [r for r in range(1, sheet.max_row + 1) if value(sheet.cell(r, 2))
              and re.fullmatch(r"\d+", value(sheet.cell(r, 2))) and not value(sheet.cell(r, 3)) is None]
    # Formula 6 has a numbered heading separate from its first component.
    if sheet.max_row >= 43 and value(sheet["B43"]) == "6":
        starts.append(43)
    starts = sorted(set(starts))
    output = []
    for start in starts:
        label = value(sheet.cell(start, 2))
        first = start if value(sheet.cell(start, 3)) else start + 1
        lines = []
        row = first
        while row <= sheet.max_row and value(sheet.cell(row, 3)):
            lines.append({"name": value(sheet.cell(row, 3)), "quantity": number(sheet.cell(row, 4)),
                          "unitPrice": number(sheet.cell(row, 6)), "lineTotal": number(sheet.cell(row, 7)),
                          "sourceRow": row})
            row += 1
        if not lines:
            raise ValueError(f"Formula {label} has no components")
        total = number(sheet.cell(row, 7))
        expected = sum(Decimal(line["lineTotal"]) for line in lines)
        if total is None or abs(Decimal(total) - expected) > Decimal("0.00001"):
            raise ValueError(f"Formula total mismatch at row {row}")
        customer = next((value(sheet.cell(i, 2)) for i in range(first, row)
                         if value(sheet.cell(i, 2)) and not value(sheet.cell(i, 2)).isdigit()), None)
        output.append({"sourceSheet": sheet.title, "sourceRows": source_rows(sheet, start, row),
                       "number": label, "customer": customer, "components": lines,
                       "total": total, "costPerUnit": number(sheet.cell(row, 8), allow_none=True)})
    return output


def september_sales(sheet):
    records = []
    for row in range(5, sheet.max_row + 1):
        for customer, date_col, quantity_col, price_col, total_col, accepted_date_col, accepted_col in [
            ("ETG", 3, 4, 5, 6, 7, 8), ("M-DAT", 11, 12, 13, 14, None, 15)
        ]:
            day = parse_date(sheet.cell(row, date_col).value)
            if day is None:
                continue
            quantity = number(sheet.cell(row, quantity_col))
            unit_price = number(sheet.cell(row, price_col))
            total = number(sheet.cell(row, total_col))
            if abs(Decimal(quantity) * Decimal(unit_price) - Decimal(total)) > Decimal("0.00001"):
                raise ValueError(f"Sales multiplication mismatch at row {row}")
            records.append({"sourceSheet": sheet.title, "sourceRows": [row], "date": day,
                            "customer": customer, "quantityM3": quantity, "unitPrice": unit_price,
                            "total": total, "acceptedDate": parse_date(sheet.cell(row, accepted_date_col).value)
                            if accepted_date_col else None,
                            "acceptedQuantity": number(sheet.cell(row, accepted_col), allow_none=True)})
    return records


def summer_metrics(sheet):
    records = []
    for row in range(1, sheet.max_row + 1):
        day = parse_date(sheet.cell(row, 3).value)
        if day is not None:
            records.append({"sourceSheet": sheet.title, "sourceRows": [row], "date": day,
                            "label": value(sheet.cell(row, 4)), "columnE": number(sheet.cell(row, 5), allow_none=True),
                            "columnF": number(sheet.cell(row, 6), allow_none=True)})
    return records


def main():
    arg = argparse.ArgumentParser()
    arg.add_argument("workbook", type=Path)
    args = arg.parse_args()
    source = args.workbook
    if source.name == "Отчет 2026 Сентябрь.xlsx":
        source_key = "2026-september"
    elif source.name == "Отчет 2026 июль-август.xlsx":
        source_key = "2026-july-august"
    else:
        raise ValueError(f"Unexpected workbook: {source.name}")
    raw = source.read_bytes()
    cached = openpyxl.load_workbook(source, data_only=True, read_only=False)
    formula = openpyxl.load_workbook(source, data_only=False, read_only=False)
    sheets = []
    for sheet in cached:
        other = formula[sheet.title]
        rows = []
        for row in sheet:
            cells = []
            for cell in row:
                paired = other[cell.coordinate]
                if cell.value is None and paired.value is None:
                    continue
                item = {"address": cell.coordinate, "value": value(cell)}
                if paired.data_type == "f":
                    item["formula"] = paired.value
                cells.append(item)
            if cells:
                rows.append({"number": row[0].row, "cells": cells})
        sheets.append({"name": sheet.title, "maxRow": sheet.max_row, "maxColumn": sheet.max_column, "rows": rows})
    production = (september_production(cached["производство"])
                  if source_key == "2026-september" else summer_production(cached["производство"]))
    sales = september_sales(cached["реализация"]) if source_key == "2026-september" else []
    metrics = summer_metrics(cached["реализация"]) if source_key != "2026-september" else []
    result = {"schemaVersion": 1, "sourceKey": source_key, "fileName": source.name,
              "sha256": hashlib.sha256(raw).hexdigest(), "sheets": sheets,
              "production": production, "formulas": formulas(cached["Формулы"]),
              "sales": sales, "unlabelledSalesMetrics": metrics}
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
